"""Excel & CSV data processing engine for church-partner-hub.

Supports:
- Dual-mode schema detection (MASTER vs SIMPLE)
- Master spreadsheet reading & writing with 100% preservation of 17 original columns & styles
- Simple address book (3 input columns -> 7 refined output columns) parsing and export (.xlsx / .csv)
- Atomic file saving & file lock detection (PermissionError handling)
"""

import csv
import os
import shutil
import tempfile
from typing import Any, Dict, List, Optional, Tuple, Union

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.core.models import ChurchRecord, SimpleAddressRecord


# 17개 고유 마스터 컬럼 표준 명칭
MASTER_17_COLUMNS = [
    "지역",
    "구분",
    "교회명",
    "담임목사",
    "교회 주소",
    "교단",
    "성도 수",
    "관리등급",
    "온도감",
    "관리유형",
    "1차 제안사업",
    "1차 제안일",
    "현재 제안단계",
    "후속 제안사업",
    "다음 액션",
    "다음 접촉일",
    "비고",
]

# 시스템 확장 메타데이터 컬럼
SYSTEM_EXT_COLUMNS = [
    "우편번호",
    "홈페이지",
    "주소검증상태",
    "홈페이지검증상태",
    "Obsidian링크",
]

# 간편 주소록 7개 표준 출력 컬럼
SIMPLE_EXPORT_COLUMNS = [
    "교회명",
    "담임목사",
    "지역",
    "도로명 주소",
    "우편번호",
    "홈페이지",
    "검증 상태",
]


class ExcelFileLockedError(PermissionError):
    """엑셀 파일이 다른 프로그램(예: Excel)에서 열려 있어 쓰기가 불가능할 때 발생하는 예외."""
    pass


class ExcelEngine:
    """엑셀 및 CSV 입출력 처리 전담 엔진 (단일 책임 원칙 준수)."""

    def __init__(self) -> None:
        self.loaded_file_path: Optional[str] = None
        self.loaded_workbook: Optional[openpyxl.Workbook] = None
        self.active_sheet_name: Optional[str] = None
        self.detected_mode: str = "MASTER"
        self.header_map: Dict[str, int] = {}  # column_name -> 1-based col_index

    # --- Mode Detection ---

    @staticmethod
    def detect_mode(file_path: str) -> str:
        """엑셀 또는 CSV 파일의 헤더를 검사하여 운영 모드(MASTER | SIMPLE)를 자동 판별합니다."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"파일을 찾을 수 없습니다: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        headers = []

        if ext == ".csv":
            # CSV 파일은 보통 간편 모드로 처리
            return "SIMPLE"

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                non_empty = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if len(non_empty) >= 2:
                    headers = [str(c).strip() if c is not None else "" for c in row]
                    break
            wb.close()
        except Exception as e:
            raise ValueError(f"엑셀 파일을 분석할 수 없습니다: {e}")

        # 마스터 특징 키워드 매칭
        master_keywords = ["구분", "관리등급", "제안사업", "온도감", "관리유형", "다음 액션", "접촉일"]
        master_matches = sum(1 for kw in master_keywords if any(kw in h for h in headers))

        # 간편 모드 필수 키워드 (담임목사, 지역, 교회명)
        has_church = any("교회" in h for h in headers)
        has_pastor = any("목사" in h or "목회자" in h for h in headers)
        has_region = any("지역" in h or "시도" in h for h in headers)

        if master_matches >= 2:
            return "MASTER"
        elif has_church and (has_pastor or has_region):
            return "SIMPLE"
        return "MASTER"  # 기본값

    # --- Master Mode I/O ---

    def load_master_records(self, file_path: str) -> List[ChurchRecord]:
        """마스터 엑셀 파일을 로드하여 서식 보존용 워크북을 캐싱하고 레코드 리스트를 생성합니다."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"파일을 찾을 수 없습니다: {file_path}")

        # 서식/수식을 보존하기 위해 data_only=False로 로드
        wb = openpyxl.load_workbook(file_path, data_only=False)
        ws = wb.active

        self.loaded_file_path = file_path
        self.loaded_workbook = wb
        self.active_sheet_name = ws.title
        self.detected_mode = "MASTER"

        # 헤더 행 탐색
        header_row_idx = 1
        header_map: Dict[str, int] = {}

        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            row_str = [str(c).strip() if c is not None else "" for c in row]
            if any("교회" in c for c in row_str):
                header_row_idx = r_idx
                for c_idx, val in enumerate(row_str, start=1):
                    if val:
                        header_map[val] = c_idx
                break

        self.header_map = header_map

        # 컬럼 인덱스 헬퍼
        def get_col(possible_names: List[str]) -> Optional[int]:
            for name in possible_names:
                for h_name, c_idx in header_map.items():
                    if name in h_name:
                        return c_idx
            return None

        col_region = get_col(["지역"])
        col_class = get_col(["구분"])
        col_name = get_col(["교회명", "교회"])
        col_pastor = get_col(["담임목사", "목사", "목회자"])
        col_addr = get_col(["주소", "도로명"])
        col_denom = get_col(["교단"])
        col_cong = get_col(["성도", "성도수", "성도 수"])
        col_tier = get_col(["관리등급", "등급"])
        col_temp = get_col(["온도감", "온도"])
        col_mgt = get_col(["관리유형", "유형"])
        col_p_camp = get_col(["1차 제안사업", "1차제안"])
        col_p_date = get_col(["1차 제안일", "1차제안일"])
        col_status = get_col(["현재 제안단계", "제안단계", "제안상태"])
        col_f_camp = get_col(["후속 제안사업", "후속제안"])
        col_n_act = get_col(["다음 액션", "다음액션", "차기액션"])
        col_n_date = get_col(["다음 접촉일", "접촉일"])
        col_remarks = get_col(["비고", "메모"])

        # 확장 컬럼
        col_zip = get_col(["우편번호"])
        col_hp = get_col(["홈페이지", "웹사이트"])
        col_v_status = get_col(["주소검증상태", "검증상태"])
        col_hp_status = get_col(["홈페이지검증상태", "홈페이지상태"])
        col_obsidian = get_col(["Obsidian링크", "옵시디언"])

        records: List[ChurchRecord] = []

        def cell_val(row_idx: int, col_idx: Optional[int]) -> str:
            if not col_idx:
                return ""
            val = ws.cell(row=row_idx, column=col_idx).value
            return str(val).strip() if val is not None else ""

        def int_val(row_idx: int, col_idx: Optional[int]) -> Optional[int]:
            if not col_idx:
                return None
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is None or val == "":
                return None
            try:
                # 숫자 또는 '1,200' 형태 처리
                clean_str = str(val).replace(",", "").strip()
                return int(float(clean_str))
            except (ValueError, TypeError):
                return None

        total_rows = ws.max_row
        for r_idx in range(header_row_idx + 1, total_rows + 1):
            c_name = cell_val(r_idx, col_name)
            if not c_name:
                continue

            rec = ChurchRecord(
                row_id=r_idx,
                region=cell_val(r_idx, col_region),
                classification=cell_val(r_idx, col_class),
                church_name=c_name,
                pastor=cell_val(r_idx, col_pastor),
                address=cell_val(r_idx, col_addr),
                denomination=cell_val(r_idx, col_denom),
                congregation_size=int_val(r_idx, col_cong),
                tier=cell_val(r_idx, col_tier),
                temperature=cell_val(r_idx, col_temp),
                management_type=cell_val(r_idx, col_mgt),
                primary_campaign=cell_val(r_idx, col_p_camp),
                primary_campaign_date=cell_val(r_idx, col_p_date),
                campaign_status=cell_val(r_idx, col_status),
                followup_campaign=cell_val(r_idx, col_f_camp),
                next_action=cell_val(r_idx, col_n_act),
                next_contact_date=cell_val(r_idx, col_n_date),
                remarks=cell_val(r_idx, col_remarks),
                zip_code=cell_val(r_idx, col_zip),
                homepage=cell_val(r_idx, col_hp),
                verification_status=cell_val(r_idx, col_v_status) or ("승인완료" if cell_val(r_idx, col_addr) else "미검증"),
                homepage_status=cell_val(r_idx, col_hp_status) or ("확인완료" if cell_val(r_idx, col_hp) else "미검증"),
                obsidian_link=cell_val(r_idx, col_obsidian),
            )
            records.append(rec)

        return records

    def save_master_records(
        self,
        records: List[ChurchRecord],
        output_path: Optional[str] = None,
    ) -> str:
        """마스터 레코드의 업데이트 사항을 원본 서식을 보존하며 안전하게(Atomic Save) 저장합니다."""
        target_path = output_path or self.loaded_file_path
        if not target_path:
            raise ValueError("저장할 파일 경로가 지정되지 않았습니다.")

        # 파일 락 검사
        self._check_file_lock(target_path)

        wb = self.loaded_workbook
        if not wb:
            # 새 워크북 생성
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "교회관리마스터"
            # 17개 헤더 + 5개 확장 헤더 작성
            all_headers = MASTER_17_COLUMNS + SYSTEM_EXT_COLUMNS
            ws.append(all_headers)
            self._apply_header_style(ws)
            for rec in records:
                row_vals = [
                    rec.region,
                    rec.classification,
                    rec.church_name,
                    rec.pastor,
                    rec.address,
                    rec.denomination,
                    rec.congregation_size,
                    rec.tier,
                    rec.temperature,
                    rec.management_type,
                    rec.primary_campaign,
                    rec.primary_campaign_date,
                    rec.campaign_status,
                    rec.followup_campaign,
                    rec.next_action,
                    rec.next_contact_date,
                    rec.remarks,
                    rec.zip_code,
                    rec.homepage,
                    rec.verification_status,
                    rec.homepage_status,
                    rec.obsidian_link,
                ]
                ws.append(row_vals)
        else:
            ws = wb[self.active_sheet_name] if self.active_sheet_name else wb.active

            # 확장 컬럼 헤더가 시트에 없으면 맨 우측에 추가
            header_row = 1
            max_col = ws.max_column
            existing_headers = {
                str(ws.cell(row=header_row, column=c).value).strip(): c
                for c in range(1, max_col + 1)
                if ws.cell(row=header_row, column=c).value is not None
            }

            for ext_col in SYSTEM_EXT_COLUMNS:
                if ext_col not in existing_headers:
                    max_col += 1
                    ws.cell(row=header_row, column=max_col, value=ext_col)
                    existing_headers[ext_col] = max_col

            # 각 레코드 데이터 갱신
            for rec in records:
                r_idx = rec.row_id
                if r_idx <= header_row:
                    continue

                # 기본 필드 매핑 및 업데이트
                field_map = {
                    "교회 주소": rec.address,
                    "주소": rec.address,
                    "우편번호": rec.zip_code,
                    "홈페이지": rec.homepage,
                    "주소검증상태": rec.verification_status,
                    "홈페이지검증상태": rec.homepage_status,
                    "Obsidian링크": rec.obsidian_link,
                }

                for col_name, val in field_map.items():
                    target_col = existing_headers.get(col_name)
                    if target_col and val:
                        ws.cell(row=r_idx, column=target_col, value=val)

        # Atomic Save 실행
        self._atomic_save(wb, target_path)
        return target_path

    # --- Simple Address Book Mode I/O ---

    def load_simple_records(self, file_path: str) -> List[SimpleAddressRecord]:
        """간편 주소록용 파일(.xlsx 또는 .csv)을 로드하여 3개 필수 컬럼을 매핑합니다."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"파일을 찾을 수 없습니다: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        records: List[SimpleAddressRecord] = []

        if ext == ".csv":
            records = self._load_simple_csv(file_path)
        else:
            records = self._load_simple_xlsx(file_path)

        self.loaded_file_path = file_path
        self.detected_mode = "SIMPLE"
        return records

    def _load_simple_csv(self, file_path: str) -> List[SimpleAddressRecord]:
        records: List[SimpleAddressRecord] = []
        encodings = ["utf-8-sig", "utf-8", "cp949", "euc-kr"]

        raw_rows = []
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    reader = csv.reader(f)
                    raw_rows = list(reader)
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if not raw_rows:
            return []

        # 헤더 매핑
        header = [c.strip() for c in raw_rows[0]]
        col_church = self._find_col_idx(header, ["교회명", "교회"])
        col_pastor = self._find_col_idx(header, ["담임목사", "목사", "목회자"])
        col_region = self._find_col_idx(header, ["지역", "시도"])
        col_addr = self._find_col_idx(header, ["주소", "도로명", "도로명주소"])
        col_zip = self._find_col_idx(header, ["우편번호", "우편"])
        col_hp = self._find_col_idx(header, ["홈페이지", "웹사이트", "url"])

        for r_idx, row in enumerate(raw_rows[1:], start=2):
            if not row:
                continue
            c_name = row[col_church].strip() if col_church is not None and col_church < len(row) else ""
            if not c_name:
                continue

            pastor = row[col_pastor].strip() if col_pastor is not None and col_pastor < len(row) else ""
            region = row[col_region].strip() if col_region is not None and col_region < len(row) else ""
            addr = row[col_addr].strip() if col_addr is not None and col_addr < len(row) else ""
            zip_code = row[col_zip].strip() if col_zip is not None and col_zip < len(row) else ""
            hp = row[col_hp].strip() if col_hp is not None and col_hp < len(row) else ""

            records.append(
                SimpleAddressRecord(
                    row_id=r_idx,
                    church_name=c_name,
                    pastor=pastor,
                    region=region,
                    road_address=addr,
                    zip_code=zip_code,
                    homepage=hp,
                    address_status="확인완료" if addr else "미검증",
                    homepage_status="확인완료" if hp else "미검증",
                )
            )
        return records

    def _load_simple_xlsx(self, file_path: str) -> List[SimpleAddressRecord]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        records: List[SimpleAddressRecord] = []

        header_row_idx = 1
        header_vals: List[str] = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [str(c).strip() if c is not None else "" for c in row]
            if any("교회" in v for v in vals):
                header_row_idx = r_idx
                header_vals = vals
                break

        col_church = self._find_col_idx(header_vals, ["교회명", "교회"])
        col_pastor = self._find_col_idx(header_vals, ["담임목사", "목사", "목회자"])
        col_region = self._find_col_idx(header_vals, ["지역", "시도"])
        col_addr = self._find_col_idx(header_vals, ["주소", "도로명", "도로명주소"])
        col_zip = self._find_col_idx(header_vals, ["우편번호", "우편"])
        col_hp = self._find_col_idx(header_vals, ["홈페이지", "웹사이트", "url"])

        for r_idx in range(header_row_idx + 1, ws.max_row + 1):
            def c_val(c_idx: Optional[int]) -> str:
                if c_idx is None:
                    return ""
                val = ws.cell(row=r_idx, column=c_idx + 1).value
                return str(val).strip() if val is not None else ""

            c_name = c_val(col_church)
            if not c_name:
                continue

            pastor = c_val(col_pastor)
            region = c_val(col_region)
            addr = c_val(col_addr)
            zip_code = c_val(col_zip)
            hp = c_val(col_hp)

            records.append(
                SimpleAddressRecord(
                    row_id=r_idx,
                    church_name=c_name,
                    pastor=pastor,
                    region=region,
                    road_address=addr,
                    zip_code=zip_code,
                    homepage=hp,
                    address_status="확인완료" if addr else "미검증",
                    homepage_status="확인완료" if hp else "미검증",
                )
            )

        wb.close()
        return records

    def export_simple_address_book(
        self,
        records: List[SimpleAddressRecord],
        output_path: str,
        file_format: str = "xlsx",
    ) -> str:
        """간편 주소록 모드 전용: 7개 정제 컬럼을 .xlsx 또는 .csv로 내보냅니다."""
        self._check_file_lock(output_path)

        if file_format.lower() == "csv" or output_path.lower().endswith(".csv"):
            return self._export_simple_csv(records, output_path)
        return self._export_simple_xlsx(records, output_path)

    def _export_simple_csv(self, records: List[SimpleAddressRecord], output_path: str) -> str:
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".csv", mode="w", encoding="utf-8-sig", newline="")
        temp_path = temp_file.name
        try:
            writer = csv.writer(temp_file)
            writer.writerow(SIMPLE_EXPORT_COLUMNS)
            for r in records:
                writer.writerow([
                    r.church_name,
                    r.pastor,
                    r.region,
                    r.road_address,
                    r.zip_code,
                    r.homepage,
                    r.address_status,
                ])
            temp_file.close()
            shutil.move(temp_path, output_path)
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
        return output_path

    def _export_simple_xlsx(self, records: List[SimpleAddressRecord], output_path: str) -> str:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "간편주소록"

        ws.append(SIMPLE_EXPORT_COLUMNS)
        self._apply_header_style(ws)

        for r in records:
            ws.append([
                r.church_name,
                r.pastor,
                r.region,
                r.road_address,
                r.zip_code,
                r.homepage,
                r.address_status,
            ])

        # 열 너비 자동 조정
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        self._atomic_save(wb, output_path)
        return output_path

    # --- Internal Utilities ---

    @staticmethod
    def _find_col_idx(headers: List[str], keywords: List[str]) -> Optional[int]:
        for kw in keywords:
            for idx, h in enumerate(headers):
                if kw in h:
                    return idx
        return None

    @staticmethod
    def _apply_header_style(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
        """기본 헤더 행에 모던한 스타일(진한 배경, 흰색 볼드 텍스트, 중앙 정렬)을 적용합니다."""
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
        alignment = Alignment(horizontal="center", vertical="center")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = alignment
        ws.row_dimensions[1].height = 28

    @staticmethod
    def _check_file_lock(file_path: str) -> None:
        """대상 파일이 다른 프로세스에 의해 잠겨 있는지 확인합니다."""
        if not os.path.exists(file_path):
            return

        try:
            with open(file_path, "r+b"):
                pass
        except PermissionError as e:
            raise ExcelFileLockedError(
                f"파일이 다른 프로그램(예: Excel)에서 열려 있어 저장할 수 없습니다.\n"
                f"파일을 닫고 다시 시도해주세요: {os.path.basename(file_path)}"
            ) from e

    @staticmethod
    def _atomic_save(wb: openpyxl.Workbook, target_path: str) -> None:
        """임시 파일에 안전하게 저장한 뒤 대상 파일로 원자적 교체(Atomic replace)를 수행합니다."""
        dir_name = os.path.dirname(os.path.abspath(target_path)) or "."
        os.makedirs(dir_name, exist_ok=True)

        temp_fd, temp_path = tempfile.mkstemp(suffix=".tmp", dir=dir_name)
        os.close(temp_fd)

        try:
            wb.save(temp_path)
            # 윈도우 환경에서 원자적 교체
            os.replace(temp_path, target_path)
        except PermissionError as e:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            raise ExcelFileLockedError(
                f"파일이 다른 프로그램에서 열려 있어 변경사항을 저장할 수 없습니다: {os.path.basename(target_path)}"
            ) from e
        except Exception:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            raise
