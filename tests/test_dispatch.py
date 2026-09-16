"""Tests for DispatchManager and Korean Chosung search utilities."""

import os
import tempfile
import openpyxl
import pytest

from src.core.dispatch import DispatchManager, extract_chosung, is_chosung_only
from src.core.models import ChurchRecord, SimpleAddressRecord


def test_extract_chosung():
    # 기본 단자음 초성 추출
    assert extract_chosung("광주") == "ㄱㅈ"
    assert extract_chosung("영락교회") == "ㅇㄹㄱㅎ"
    assert extract_chosung("겨자씨") == "ㄱㅈㅆ"
    # 쌍자음 정규화
    assert extract_chosung("겨자씨", normalize_complex=True) == "ㄱㅈㅅ"
    # 영문 및 숫자 보존
    assert extract_chosung("Church 123") == "Church 123"


def test_is_chosung_only():
    assert is_chosung_only("ㄱㅈㅅ") is True
    assert is_chosung_only("ㄱ ㅈ ㅅ") is True
    assert is_chosung_only("광주") is False
    assert is_chosung_only("ㄱㅈa") is False
    assert is_chosung_only("") is False


def test_dispatch_search():
    records = [
        ChurchRecord(
            row_id=1,
            region="광주",
            church_name="광주겨자씨교회",
            pastor="나학수",
            address="광주 남구 봉선로 12",
        ),
        ChurchRecord(
            row_id=2,
            region="광주",
            church_name="광주동성교회",
            pastor="안성주",
            address="광주 동구 필문대로 205",
        ),
        ChurchRecord(
            row_id=3,
            region="서울",
            church_name="영락교회",
            pastor="김운성",
            address="서울 중구 수표로 33",
        ),
    ]

    # 1. 초성 검색: 'ㄱㅈㅅ' -> 광주겨자씨교회
    res = DispatchManager.search(records, "ㄱㅈㅅ")
    assert len(res) == 1
    assert res[0].church_name == "광주겨자씨교회"

    # 2. 초성 검색: 'ㅇㄹ' -> 영락교회
    res = DispatchManager.search(records, "ㅇㄹ")
    assert len(res) == 1
    assert res[0].church_name == "영락교회"

    # 3. 일반 검색 (목회자명)
    res = DispatchManager.search(records, "안성주")
    assert len(res) == 1
    assert res[0].church_name == "광주동성교회"

    # 4. 일반 검색 (지역)
    res = DispatchManager.search(records, "광주")
    assert len(res) == 2

    # 5. 빈 검색어
    res = DispatchManager.search(records, "")
    assert len(res) == 3


def test_dispatch_formatting():
    record = ChurchRecord(
        row_id=1,
        region="광주",
        church_name="광주겨자씨교회",
        pastor="나학수",
        address="광주 남구 봉선로 12",
        zip_code="61642",
        homepage="http://mustardseed.or.kr",
        remarks="기념품 발송",
    )

    official = DispatchManager.format_official(record)
    assert "수신: 광주겨자씨교회 (나학수 목사 귀하)" in official
    assert "(61642)" in official
    assert "광주 남구 봉선로 12" in official
    assert "비고: 기념품 발송" in official

    parcel = DispatchManager.format_parcel(record)
    assert "받는분: 나학수 목사 (광주겨자씨교회)" in parcel
    assert "우편번호: 61642" in parcel
    assert "배송주소: 광주 남구 봉선로 12" in parcel

    tsv = DispatchManager.format_tsv(record)
    parts = tsv.split("\t")
    assert parts[0] == "광주겨자씨교회"
    assert parts[1] == "나학수"
    assert parts[2] == "광주"
    assert parts[3] == "61642"


def test_export_dispatch_list():
    records = [
        ChurchRecord(
            row_id=1,
            region="광주",
            church_name="광주겨자씨교회",
            pastor="나학수",
            address="광주 남구 봉선로 12",
            zip_code="61642",
        ),
    ]

    with tempfile.TemporaryDirectory() as tmpdir:
        out_file = os.path.join(tmpdir, "dispatch.xlsx")
        DispatchManager.export_dispatch_list(records, out_file)

        assert os.path.exists(out_file)
        wb = openpyxl.load_workbook(out_file)
        ws = wb.active
        assert ws.title == "발송명단"
        assert ws.cell(row=1, column=2).value == "교회명"
        assert ws.cell(row=2, column=2).value == "광주겨자씨교회"
        wb.close()
