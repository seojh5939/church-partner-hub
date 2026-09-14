"""Tests for ExcelEngine: dual-mode schema, master style preservation, simple export, and atomic save."""

import csv
import os
import tempfile
import openpyxl
import pytest

from src.core.excel_engine import (
    ExcelEngine,
    ExcelFileLockedError,
    MASTER_17_COLUMNS,
    SIMPLE_EXPORT_COLUMNS,
)
from src.core.models import ChurchRecord, SimpleAddressRecord


@pytest.fixture
def sample_master_excel():
    """17개 마스터 컬럼을 포함하는 임시 엑셀 파일을 생성합니다."""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(MASTER_17_COLUMNS)
    ws.append([
        "광주", "이사교회", "광주겨자씨교회", "나학수", "광주 남구 봉선로 12", "예장합동",
        3500, "A", "Hot", "본부집중", "희망친구 결연", "2026-03-15", "제안완료",
        "국내위기가정지원", "방문 미팅", "2026-09-25", "창립기념 선물",
    ])
    ws.append([
        "광주", "타겟교회", "광주동성교회", "안성주", "", "예장통합",
        800, "B", "Warm", "지역본부", "긴급구호", "2026-05-10", "검토중",
        "우물파기", "자료 이메일", "2026-09-20", "",
    ])
    wb.save(tmp.name)
    wb.close()

    yield tmp.name
    if os.path.exists(tmp.name):
        os.remove(tmp.name)


@pytest.fixture
def sample_simple_excel():
    """3개 필수 컬럼[담임목사, 지역, 교회명]을 포함하는 임시 간편 엑셀 파일을 생성합니다."""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["담임목사", "지역", "교회명"])
    ws.append(["나학수", "광주", "광주겨자씨교회"])
    ws.append(["안성주", "광주", "광주동성교회"])
    wb.save(tmp.name)
    wb.close()

    yield tmp.name
    if os.path.exists(tmp.name):
        os.remove(tmp.name)


@pytest.fixture
def sample_simple_csv():
    """3개 필수 컬럼을 포함하는 임시 CSV 파일을 생성합니다."""
    tmp = tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w", encoding="utf-8-sig", newline="")
    writer = csv.writer(tmp)
    writer.writerow(["담임목사", "지역", "교회명"])
    writer.writerow(["나학수", "광주", "광주겨자씨교회"])
    tmp.close()

    yield tmp.name
    if os.path.exists(tmp.name):
        os.remove(tmp.name)


def test_detect_mode(sample_master_excel, sample_simple_excel, sample_simple_csv):
    engine = ExcelEngine()
    assert engine.detect_mode(sample_master_excel) == "MASTER"
    assert engine.detect_mode(sample_simple_excel) == "SIMPLE"
    assert engine.detect_mode(sample_simple_csv) == "SIMPLE"


def test_load_and_save_master_records(sample_master_excel):
    engine = ExcelEngine()
    records = engine.load_master_records(sample_master_excel)

    assert len(records) == 2
    assert records[0].church_name == "광주겨자씨교회"
    assert records[0].pastor == "나학수"
    assert records[0].congregation_size == 3500
    assert records[0].verification_status == "승인완료"
    assert records[1].church_name == "광주동성교회"
    assert records[1].verification_status == "미검증"

    # 주소 및 우편번호 업데이트 후 저장
    records[1].address = "광주광역시 동구 필문대로 205"
    records[1].zip_code = "61448"
    records[1].verification_status = "승인완료"

    saved_path = engine.save_master_records(records)
    assert saved_path == sample_master_excel

    # 다시 로드하여 보존 및 업데이트 여부 확인
    reloaded_engine = ExcelEngine()
    reloaded = reloaded_engine.load_master_records(saved_path)
    assert len(reloaded) == 2
    assert reloaded[1].address == "광주광역시 동구 필문대로 205"
    assert reloaded[1].zip_code == "61448"
    assert reloaded[1].verification_status == "승인완료"


def test_load_and_export_simple_records(sample_simple_excel, sample_simple_csv):
    engine = ExcelEngine()

    # 1. XLSX 로드
    records_xlsx = engine.load_simple_records(sample_simple_excel)
    assert len(records_xlsx) == 2
    assert records_xlsx[0].church_name == "광주겨자씨교회"
    assert records_xlsx[0].pastor == "나학수"
    assert records_xlsx[0].region == "광주"

    # 2. CSV 로드
    records_csv = engine.load_simple_records(sample_simple_csv)
    assert len(records_csv) == 1
    assert records_csv[0].church_name == "광주겨자씨교회"

    # 3. 7개 컬럼 XLSX 내보내기
    with tempfile.TemporaryDirectory() as tmpdir:
        out_xlsx = os.path.join(tmpdir, "exported.xlsx")
        records_xlsx[0].road_address = "광주광역시 남구 봉선로 12"
        records_xlsx[0].zip_code = "61642"
        records_xlsx[0].homepage = "http://mustardseed.or.kr"
        records_xlsx[0].address_status = "확인완료"

        engine.export_simple_address_book(records_xlsx, out_xlsx, file_format="xlsx")
        assert os.path.exists(out_xlsx)

        wb = openpyxl.load_workbook(out_xlsx)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, 8)]
        assert headers == SIMPLE_EXPORT_COLUMNS
        assert ws.cell(row=2, column=1).value == "광주겨자씨교회"
        assert ws.cell(row=2, column=4).value == "광주광역시 남구 봉선로 12"
        wb.close()

        # 4. 7개 컬럼 CSV 내보내기
        out_csv = os.path.join(tmpdir, "exported.csv")
        engine.export_simple_address_book(records_xlsx, out_csv, file_format="csv")
        assert os.path.exists(out_csv)

        with open(out_csv, "r", encoding="utf-8-sig") as f:
            reader = list(csv.reader(f))
            assert reader[0] == SIMPLE_EXPORT_COLUMNS
            assert reader[1][0] == "광주겨자씨교회"
            assert reader[1][3] == "광주광역시 남구 봉선로 12"


def test_file_lock_detection(sample_master_excel):
    engine = ExcelEngine()
    records = engine.load_master_records(sample_master_excel)

    # 파일을 배타적 잠금(exclusive lock) 모드로 열어둠
    with open(sample_master_excel, "r+b") as locked_file:
        # 윈도우 msvcrt lock 적용 시도
        try:
            import msvcrt
            msvcrt.locking(locked_file.fileno(), msvcrt.LK_NBLCK, 1)
            is_locked = True
        except Exception:
            is_locked = False

        if is_locked:
            with pytest.raises(ExcelFileLockedError):
                engine.save_master_records(records, sample_master_excel)
            # 락 해제
            msvcrt.locking(locked_file.fileno(), msvcrt.LK_UNLCK, 1)
